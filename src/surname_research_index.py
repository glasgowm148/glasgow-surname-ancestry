#!/usr/bin/env python3
"""Build navigable surname-wide indexes from the workbook and research cases."""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook


WIKITREE_BASE = "https://www.wikitree.com/wiki/"
PROFILE_INDEX_JSON = "profiles.json"
RESEARCHED_INDEX_MD = "researched-profiles.md"
EARLY_INDEX_MD = "early-profiles.md"
EARLY_IRISH_INDEX_MD = "early-irish-profiles.md"
FATHER_ASSESSMENTS_JSON = "father-assessments.json"
IRISH_COHORT_AUDIT_JSON = "irish-cohort-audit.json"

FATHER_STATUS_LABELS = {
    "confirmed": "Confirmed by record",
    "plausible": "Plausible",
    "uncertain": "Uncertain",
    "contradicted": "Contradicted / unsupported",
}

AREA_ORDER = (
    "Scotland: Lanarkshire and Glasgow",
    "Scotland: Ayrshire",
    "Scotland: Edinburgh and the Lothians",
    "Scotland: Borders",
    "Scotland: Stirlingshire",
    "Scotland: Argyll",
    "Scotland: other or unspecified",
    "Ireland: County Antrim",
    "Ireland: County Londonderry",
    "Ireland: County Tyrone",
    "Ireland: other or unspecified",
    "England",
    "North America and Caribbean",
    "Other or location unknown",
)

IRISH_AREA_ORDER = (
    "Ireland: County Antrim",
    "Ireland: County Londonderry",
    "Ireland: County Tyrone",
    "Ireland: County Donegal",
    "Ireland: County Armagh",
    "Ireland: County Down",
    "Ireland: County Fermanagh",
    "Ireland: County Leitrim",
    "Ireland: County Galway",
    "Ireland: County Dublin",
    "Ireland: County Cork",
    "Ireland: Queen's County / Laois",
    "Ireland: Belfast, county unresolved",
    "Ireland: other or unspecified",
)


def text(value: Any, default: str = "unknown") -> str:
    if value in (None, ""):
        return default
    return str(value).strip() or default


def clean_date(value: Any) -> str:
    value = text(value)
    if re.fullmatch(r"0000(?:-00(?:-00)?)?", value):
        return "unknown"
    if re.fullmatch(r"\d{4}-00-00", value):
        return value[:4]
    if re.fullmatch(r"\d{4}-\d{2}-00", value):
        return value[:7]
    return value


def birth_year(profile: dict[str, Any]) -> int | None:
    match = re.search(
        r"(?<!\d)(1[2-9]\d{2}|20\d{2})(?!\d)",
        clean_date(profile.get("birth_date")),
    )
    return int(match.group(1)) if match else None


def identity_label(profile: dict[str, Any]) -> str:
    birth_date = clean_date(profile.get("birth_date"))
    birth_location = text(profile.get("birth_location"))
    birth = "birth unknown" if birth_date == "unknown" else f"born {birth_date}"
    location = (
        "birthplace unknown" if birth_location == "unknown" else birth_location
    )
    return (
        f"{text(profile.get('display_name'), text(profile.get('wikitree_id')))} "
        f"({text(profile.get('wikitree_id'))}), {birth}, {location}"
    )


def profile_link(profile: dict[str, Any]) -> str:
    wt_id = text(profile.get("wikitree_id"))
    return f"[{identity_label(profile)}]({WIKITREE_BASE}{wt_id})"


def concise_profile_link(profile: dict[str, Any]) -> str:
    wt_id = text(profile.get("wikitree_id"))
    name = text(profile.get("display_name"), wt_id)
    year = birth_year(profile)
    lifespan = str(year) if year is not None else "birth unknown"
    return f"[{name} ({lifespan})]({WIKITREE_BASE}{wt_id})"


def load_workbook_profiles(path: Path) -> dict[str, dict[str, Any]]:
    profiles: dict[str, dict[str, Any]] = {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["TableData"]
    rows = sheet.iter_rows(values_only=True)
    headers = [text(value, "") for value in next(rows)]

    for values in rows:
        row = dict(zip(headers, values))
        wt_id = text(row.get("WT ID"), "")
        if not re.fullmatch(r"Glasgow-\d+", wt_id):
            continue
        first = text(row.get("First"), "")
        surname = text(row.get("LNAB"), text(row.get("Current"), ""))
        profiles[wt_id] = {
            "wikitree_id": wt_id,
            "display_name": " ".join(part for part in (first, surname) if part) or wt_id,
            "birth_date": row.get("Birth Date"),
            "birth_location": row.get("Birth Place"),
            "death_date": row.get("Death Date"),
            "death_location": row.get("Death Place"),
            "relations": {},
            "research_cases": [],
            "direct_case": False,
            "early_evidence": False,
            "workbook_modified": row.get("Modified"),
        }
    workbook.close()
    return profiles


def merge_case_indexes(
    profiles: dict[str, dict[str, Any]],
    research_root: Path,
) -> None:
    case_dirs = sorted(path for path in research_root.glob("Glasgow-*") if path.is_dir())
    for case_dir in case_dirs:
        case_id = case_dir.name
        case_profile = profiles.setdefault(
            case_id,
            {
                "wikitree_id": case_id,
                "display_name": case_id,
                "birth_date": None,
                "birth_location": None,
                "death_date": None,
                "death_location": None,
                "relations": {},
                "research_cases": [],
                "direct_case": True,
                "early_evidence": False,
            },
        )
        case_profile["direct_case"] = True
        case_profile["case_files"] = sorted(
            path.name for path in case_dir.iterdir() if path.is_file()
        )

        index_path = case_dir / "evidence_index.json"
        if not index_path.exists():
            continue
        index = json.loads(index_path.read_text(encoding="utf-8"))
        for wt_id, summary in index.get("profiles", {}).items():
            if not isinstance(summary, dict) or not re.fullmatch(r"Glasgow-\d+", wt_id):
                continue
            profile = profiles.setdefault(wt_id, {"wikitree_id": wt_id})
            for key in (
                "display_name",
                "birth_date",
                "birth_location",
                "death_date",
                "death_location",
            ):
                if summary.get(key) not in (None, ""):
                    profile[key] = summary[key]
            profile.setdefault("research_cases", [])
            if case_id not in profile["research_cases"]:
                profile["research_cases"].append(case_id)
            profile["direct_case"] = (research_root / wt_id).is_dir()
            profile.setdefault("early_evidence", False)
            relations = profile.setdefault("relations", {})
            for relation, ids in summary.get("relations", {}).items():
                relations[relation] = sorted(
                    set(relations.get(relation, []))
                    | {str(value) for value in ids if str(value).startswith("Glasgow-")}
                )


def add_markdown_profiles(
    profiles: dict[str, dict[str, Any]],
    path: Path,
    early: bool = False,
) -> None:
    if not path.exists():
        return
    content = path.read_text(encoding="utf-8")
    ids = set(re.findall(r"Glasgow-\d+", content))
    for wt_id in ids:
        profile = profiles.setdefault(
            wt_id,
            {
                "wikitree_id": wt_id,
                "display_name": wt_id,
                "birth_date": None,
                "birth_location": None,
                "death_date": None,
                "death_location": None,
                "relations": {},
                "research_cases": [],
                "direct_case": False,
            },
        )
        profile.setdefault("research_cases", [])
        profile.setdefault("relations", {})
        profile["early_evidence"] = bool(profile.get("early_evidence") or early)


def research_links(profile: dict[str, Any]) -> str:
    wt_id = text(profile.get("wikitree_id"))
    links: list[str] = []
    if profile.get("direct_case"):
        base = f"../../research/{wt_id}"
        files = set(profile.get("case_files", []))
        if "findings.md" in files:
            links.append(f"[findings]({base}/findings.md)")
        if f"{wt_id}.md" in files:
            links.append(f"[WikiTree queue]({base}/{wt_id}.md)")
        if "research_brief.md" in files:
            links.append(f"[brief]({base}/research_brief.md)")
    other_cases = [case for case in profile.get("research_cases", []) if case != wt_id]
    if other_cases:
        links.append(
            "mentioned in "
            + ", ".join(
                f"[{case}](../../research/{case}/research_brief.md)" for case in other_cases
            )
        )
    return "; ".join(links) or "Not yet researched in a case folder"


def findings_link(profile: dict[str, Any]) -> str:
    wt_id = text(profile.get("wikitree_id"))
    if profile.get("direct_case") and "findings.md" in set(
        profile.get("case_files", [])
    ):
        return f"[findings](../../research/{wt_id}/findings.md)"
    return "Not yet researched"


def relation_links(profile: dict[str, Any], profiles: dict[str, dict[str, Any]]) -> str:
    groups: list[str] = []
    for relation in ("parents", "siblings", "spouses", "children"):
        ids = profile.get("relations", {}).get(relation, [])
        if not ids:
            continue
        labels = []
        for wt_id in ids:
            relative = profiles.get(wt_id, {"wikitree_id": wt_id})
            labels.append(profile_link(relative))
        groups.append(f"{relation}: " + ", ".join(labels))
    return "<br>".join(groups) or "None captured"


def load_father_assessments(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    assessments = value.get("assessments", {})
    if not isinstance(assessments, dict):
        raise ValueError(f"Invalid assessments object in {path}")
    for wt_id, assessment in assessments.items():
        if not isinstance(assessment, dict):
            raise ValueError(f"Invalid father assessment for {wt_id}")
        status = assessment.get("status")
        if status not in FATHER_STATUS_LABELS:
            raise ValueError(f"Invalid father status {status!r} for {wt_id}")
    return assessments


def load_irish_cohort_audit(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    assessments = value.get("assessments", {})
    if not isinstance(assessments, dict):
        raise ValueError(f"Invalid assessments object in {path}")
    for wt_id, assessment in assessments.items():
        if not isinstance(assessment, dict):
            raise ValueError(f"Invalid Irish cohort assessment for {wt_id}")
        if assessment.get("action") not in {"retain", "reclassify", "exclude"}:
            raise ValueError(
                f"Invalid Irish cohort action {assessment.get('action')!r} for {wt_id}"
            )
    return assessments


def area_for_location(value: Any) -> str:
    place = text(value, "").lower()
    if not place:
        return "Other or location unknown"

    if "antrim" in place:
        return "Ireland: County Antrim"
    if "londonderry" in place or "county derry" in place:
        return "Ireland: County Londonderry"
    if "tyrone" in place:
        return "Ireland: County Tyrone"
    if "donegal" in place:
        return "Ireland: County Donegal"
    if "armagh" in place:
        return "Ireland: County Armagh"
    if re.search(r"\bcounty down\b|\bco\.? down\b", place):
        return "Ireland: County Down"
    if "fermanagh" in place:
        return "Ireland: County Fermanagh"
    if "leitrim" in place:
        return "Ireland: County Leitrim"
    if "galway" in place:
        return "Ireland: County Galway"
    if "dublin" in place:
        return "Ireland: County Dublin"
    if "cork" in place:
        return "Ireland: County Cork"
    if "queen's county" in place or "queens county" in place or "laois" in place:
        return "Ireland: Queen's County / Laois"
    if "belfast" in place:
        return "Ireland: Belfast, county unresolved"
    if "ireland" in place or "ulster" in place:
        return "Ireland: other or unspecified"

    if any(
        value in place
        for value in (
            "ayrshire",
            "ayr",
            "irvine",
            "kilwinning",
            "kilbirnie",
            "stevenston",
        )
    ):
        return "Scotland: Ayrshire"
    if any(
        value in place
        for value in (
            "lanarkshire",
            "lanark",
            "glasgow",
            "biggar",
            "quothquan",
            "libberton",
        )
    ):
        return "Scotland: Lanarkshire and Glasgow"
    if any(
        value in place
        for value in (
            "midlothian",
            "edinburgh",
            "lothian",
            "linlithgow",
            "inveresk",
            "musselburgh",
            "mid calder",
            "kirknewton",
            "duddingston",
            "leith",
        )
    ):
        return "Scotland: Edinburgh and the Lothians"
    if any(
        value in place
        for value in (
            "berwickshire",
            "roxburghshire",
            "peeblesshire",
            "selkirkshire",
            "cavers",
            "hawick",
            "jedburgh",
            "polwarth",
        )
    ):
        return "Scotland: Borders"
    if "stirling" in place or "strathblane" in place:
        return "Scotland: Stirlingshire"
    if "argyll" in place or "campbeltown" in place:
        return "Scotland: Argyll"
    if "scotland" in place:
        return "Scotland: other or unspecified"

    if "england" in place or any(
        value in place
        for value in ("northumberland", "lancashire", "norfolk", "somerset")
    ):
        return "England"
    if any(
        value in place
        for value in (
            "united states",
            "america",
            "virginia",
            "maryland",
            "carolina",
            "barbados",
        )
    ):
        return "North America and Caribbean"
    return "Other or location unknown"


def area_group(profile: dict[str, Any]) -> str:
    birth_area = area_for_location(profile.get("birth_location"))
    death_area = area_for_location(profile.get("death_location"))
    if birth_area == "Other or location unknown":
        return death_area
    if (
        birth_area == "Scotland: other or unspecified"
        and death_area.startswith("Scotland:")
        and death_area != birth_area
    ):
        return death_area
    if (
        birth_area == "Ireland: other or unspecified"
        and death_area.startswith("Ireland:")
        and death_area != birth_area
    ):
        return death_area
    return birth_area


def irish_area_group(profile: dict[str, Any]) -> str | None:
    birth_area = area_for_location(profile.get("birth_location"))
    if birth_area.startswith("Ireland:"):
        return birth_area
    if birth_area == "Other or location unknown":
        death_area = area_for_location(profile.get("death_location"))
        if death_area.startswith("Ireland:"):
            return death_area
    return None


def render_researched_profiles(profiles: dict[str, dict[str, Any]]) -> str:
    selected = [profile for profile in profiles.values() if profile.get("direct_case")]
    selected.sort(key=lambda profile: (birth_year(profile) or 9999, identity_label(profile)))
    lines = [
        "# Researched Glasgow profiles",
        "",
        "> Generated from `research/Glasgow-*/evidence_index.json`. Tree links are context, not relationship proof.",
        "",
        "| Person | Case navigation | Captured WikiTree links |",
        "| --- | --- | --- |",
    ]
    for profile in selected:
        lines.append(
            f"| {profile_link(profile)} | {research_links(profile)} | "
            f"{relation_links(profile, profiles)} |"
        )
    return "\n".join(lines) + "\n"


def early_band(profile: dict[str, Any]) -> str:
    year = birth_year(profile)
    if year is None:
        return "Early evidence, birth year unresolved"
    if year < 1500:
        return "Before 1500"
    if year < 1600:
        return "16th century (1500-1599): highest priority"
    if year < 1700:
        return "17th century (1600-1699)"
    return "1700 boundary year"


def father_columns(
    profile: dict[str, Any],
    profiles: dict[str, dict[str, Any]],
    assessments: dict[str, dict[str, Any]],
) -> tuple[str, str, str]:
    wt_id = text(profile.get("wikitree_id"))
    assessment = assessments.get(wt_id)
    captured_parents = profile.get("relations", {}).get("parents", [])

    if assessment:
        father_id = text(assessment.get("father_id"), "")
        father_label = text(assessment.get("father_label"), "")
        father = (
            concise_profile_link(profiles.get(father_id, {"wikitree_id": father_id}))
            if father_id
            else father_label or "Father not identified"
        )
        status = FATHER_STATUS_LABELS[str(assessment["status"])]
        basis = text(assessment.get("basis"), "No assessment basis recorded")
        return father, status, basis

    if captured_parents:
        father_links = [
            concise_profile_link(profiles.get(parent_id, {"wikitree_id": parent_id}))
            for parent_id in captured_parents
        ]
        return (
            ", ".join(father_links),
            "Uncertain",
            "Attached in captured WikiTree data; not evidence-audited.",
        )

    return (
        "Not captured",
        "Not assessed",
        "The workbook has no parent fields and no local case capture supplied a father.",
    )


def father_status_key(
    profile: dict[str, Any],
    assessments: dict[str, dict[str, Any]],
) -> str:
    assessment = assessments.get(text(profile.get("wikitree_id")))
    if assessment:
        return str(assessment["status"])
    if profile.get("relations", {}).get("parents"):
        return "uncertain"
    return "not_assessed"


def irish_band(profile: dict[str, Any]) -> str:
    year = birth_year(profile)
    if year is None:
        return "Birth year unresolved"
    if year < 1600:
        return "Before 1600"
    if year < 1700:
        return "17th century (1600-1699)"
    return "18th century (1700-1799)"


def render_profile_sections(
    lines: list[str],
    selected: list[dict[str, Any]],
    profiles: dict[str, dict[str, Any]],
    father_assessments: dict[str, dict[str, Any]],
    bands: tuple[str, ...],
    band_grouper: Callable[[dict[str, Any]], str],
    areas: tuple[str, ...],
    area_grouper: Callable[[dict[str, Any]], str | None],
) -> None:
    selected.sort(key=lambda profile: (birth_year(profile) or 9999, identity_label(profile)))
    for band in bands:
        band_profiles = [profile for profile in selected if band_grouper(profile) == band]
        if not band_profiles:
            continue
        lines.extend(
            [
                f"## {band} ({len(band_profiles)})",
                "",
                "| Location | Profiles | Share | Confirmed | Plausible | Uncertain | Contradicted / unsupported | Not assessed |",
                "| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
            ]
        )
        for area in areas:
            area_profiles = [
                profile for profile in band_profiles if area_grouper(profile) == area
            ]
            area_count = len(area_profiles)
            if area_count:
                share = area_count / len(band_profiles) * 100
                status_counts = {
                    status: sum(
                        father_status_key(profile, father_assessments) == status
                        for profile in area_profiles
                    )
                    for status in (
                        "confirmed",
                        "plausible",
                        "uncertain",
                        "contradicted",
                        "not_assessed",
                    )
                }
                lines.append(
                    f"| {area} | {area_count} | {share:.1f}% | "
                    f"{status_counts['confirmed']} | {status_counts['plausible']} | "
                    f"{status_counts['uncertain']} | {status_counts['contradicted']} | "
                    f"{status_counts['not_assessed']} |"
                )
        total_status_counts = {
            status: sum(
                father_status_key(profile, father_assessments) == status
                for profile in band_profiles
            )
            for status in (
                "confirmed",
                "plausible",
                "uncertain",
                "contradicted",
                "not_assessed",
            )
        }
        lines.append(
            f"| **Total** | **{len(band_profiles)}** | **100.0%** | "
            f"**{total_status_counts['confirmed']}** | "
            f"**{total_status_counts['plausible']}** | "
            f"**{total_status_counts['uncertain']}** | "
            f"**{total_status_counts['contradicted']}** | "
            f"**{total_status_counts['not_assessed']}** |"
        )
        lines.append("")
        for area in areas:
            area_profiles = [
                profile for profile in band_profiles if area_grouper(profile) == area
            ]
            if not area_profiles:
                continue
            lines.extend(
                [
                    f"### {area} ({len(area_profiles)})",
                    "",
                    "| Person | Father / assessment | Assessment basis | Findings |",
                    "| --- | --- | --- | --- |",
                ]
            )
            for profile in area_profiles:
                father, status, basis = father_columns(
                    profile,
                    profiles,
                    father_assessments,
                )
                lines.append(
                    f"| {concise_profile_link(profile)} | {father}<br>**{status}** | "
                    f"{basis} | {findings_link(profile)} |"
                )
            lines.append("")


def render_early_profiles(
    profiles: dict[str, dict[str, Any]],
    father_assessments: dict[str, dict[str, Any]],
) -> str:
    selected = [
        profile
        for profile in profiles.values()
        if (birth_year(profile) is not None and birth_year(profile) <= 1700)
        or profile.get("early_evidence")
    ]
    lines = [
        "# Glasgow profiles and occurrences before 1701",
        "",
        f"**Total indexed: {len(selected)} profiles and early-record references.**",
        "",
        "> Generated navigation index. Areas use a specific birthplace when recorded, otherwise a more specific death place; this is navigation, not migration or kinship proof. An attached WikiTree father defaults to uncertain until the manual father-assessment register cites evidence; missing local parent data is shown as not assessed, not as proof that the father is unknown.",
        "",
        "Father status: **Confirmed by record** = a relational record plus a resolved identity; **Plausible** = a conflict-free indirect case; **Uncertain** = attached or proposed without sufficient proof; **Contradicted / unsupported** = reviewed evidence conflicts with or does not support the attachment.",
        "",
    ]
    render_profile_sections(
        lines,
        selected,
        profiles,
        father_assessments,
        (
            "Before 1500",
            "16th century (1500-1599): highest priority",
            "17th century (1600-1699)",
            "1700 boundary year",
            "Early evidence, birth year unresolved",
        ),
        early_band,
        AREA_ORDER,
        area_group,
    )
    return "\n".join(lines).rstrip() + "\n"


def render_early_irish_profiles(
    profiles: dict[str, dict[str, Any]],
    father_assessments: dict[str, dict[str, Any]],
    cohort_audit: dict[str, dict[str, Any]] | None = None,
) -> str:
    cohort_audit = cohort_audit or {}
    snapshot_selected = [
        profile
        for profile in profiles.values()
        if birth_year(profile) is not None
        and birth_year(profile) <= 1799
        and irish_area_group(profile) is not None
    ]
    selected = []
    for profile in snapshot_selected:
        assessment = cohort_audit.get(text(profile.get("wikitree_id")), {})
        if assessment.get("action") == "exclude":
            continue
        working = dict(profile)
        if "working_birth_date" in assessment:
            working["birth_date"] = assessment["working_birth_date"]
        if assessment.get("working_display_name"):
            working["display_name"] = assessment["working_display_name"]
        if assessment.get("working_band"):
            working["_irish_band"] = assessment["working_band"]
        selected.append(working)
    count_line = (
        f"**Snapshot candidates: {len(snapshot_selected)} profiles. "
        f"Evidence-audited working cohort: {len(selected)}.**"
        if cohort_audit
        else f"**Total indexed: {len(selected)} profiles.**"
    )
    lines = [
        "# Irish Glasgow profiles before 1800",
        "",
        count_line,
        "",
        "> Generated navigation index. Snapshot candidates have an Irish birthplace, or an Irish death/residence when birthplace is missing. Manual cohort audits may correct the working century or exclude a profile whose supposed Irish event is unsupported; these overrides do not alter the workbook or WikiTree. County groupings follow the recorded place and do not prove origin, migration, or kinship.",
        "",
        "Father status: **Confirmed by record** = a relational record plus a resolved identity; **Plausible** = a conflict-free indirect case; **Uncertain** = attached or proposed without sufficient proof; **Contradicted / unsupported** = reviewed evidence conflicts with or does not support the attachment.",
        "",
    ]
    audited_rows = [
        (wt_id, assessment)
        for wt_id, assessment in cohort_audit.items()
        if wt_id in profiles
    ]
    if audited_rows:
        lines.extend(
            [
                "## Cohort audit overrides",
                "",
                "| Person | Working treatment | Basis | Findings |",
                "| --- | --- | --- | --- |",
            ]
        )
        action_labels = {
            "retain": "Retain with corrected evidence date",
            "reclassify": "Reclassify century",
            "exclude": "Exclude pending Irish evidence",
        }
        for wt_id, assessment in sorted(audited_rows):
            profile = dict(profiles[wt_id])
            if "working_birth_date" in assessment:
                profile["birth_date"] = assessment["working_birth_date"]
            if assessment.get("working_display_name"):
                profile["display_name"] = assessment["working_display_name"]
            lines.append(
                f"| {concise_profile_link(profile)} | "
                f"{action_labels[str(assessment['action'])]} | "
                f"{text(assessment.get('basis'))} | {findings_link(profile)} |"
            )
        lines.append("")
    render_profile_sections(
        lines,
        selected,
        profiles,
        father_assessments,
        (
            "Before 1600",
            "17th century (1600-1699)",
            "18th century (1700-1799)",
        ),
        lambda profile: text(profile.get("_irish_band"), irish_band(profile)),
        IRISH_AREA_ORDER,
        irish_area_group,
    )
    return "\n".join(lines).rstrip() + "\n"


def serialisable_profiles(profiles: dict[str, dict[str, Any]]) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "profile_count": len(profiles),
        "profiles": dict(sorted(profiles.items())),
    }


def rebuild_surname_indexes(workspace_root: Path) -> list[Path]:
    workspace_root = workspace_root.resolve()
    surname_root = workspace_root / "surname-research"
    index_root = surname_root / "indexes"
    index_root.mkdir(parents=True, exist_ok=True)
    workbook_path = workspace_root / "data" / "reference" / "glasgow-one-tree.xlsx"
    if not workbook_path.exists():
        # Compatibility for callers that construct the former flat fixture layout.
        workbook_path = workspace_root / "glasgow-one-tree.xlsx"
    profiles = load_workbook_profiles(workbook_path)
    merge_case_indexes(profiles, workspace_root / "research")
    add_markdown_profiles(profiles, surname_root / "clusters" / "README.md")
    add_markdown_profiles(
        profiles,
        surname_root / "pre-1700" / "evidence-register.md",
        early=True,
    )
    father_assessments = load_father_assessments(
        index_root / FATHER_ASSESSMENTS_JSON
    )
    irish_cohort_audit = load_irish_cohort_audit(
        index_root / IRISH_COHORT_AUDIT_JSON
    )

    outputs = [
        index_root / PROFILE_INDEX_JSON,
        index_root / RESEARCHED_INDEX_MD,
        index_root / EARLY_INDEX_MD,
        index_root / EARLY_IRISH_INDEX_MD,
    ]
    outputs[0].write_text(
        json.dumps(
            serialisable_profiles(profiles),
            ensure_ascii=False,
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )
    outputs[1].write_text(render_researched_profiles(profiles), encoding="utf-8")
    outputs[2].write_text(
        render_early_profiles(profiles, father_assessments),
        encoding="utf-8",
    )
    outputs[3].write_text(
        render_early_irish_profiles(
            profiles,
            father_assessments,
            irish_cohort_audit,
        ),
        encoding="utf-8",
    )
    return outputs


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workspace", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    outputs = rebuild_surname_indexes(args.workspace)
    for output in outputs:
        print(f"Wrote {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
